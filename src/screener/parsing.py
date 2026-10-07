"""Extract text and hyperlinks from a resume file.

PDF is the required format. PyMuPDF is the primary parser because it returns
text in content order (multi-column resumes stay readable) and exposes link
annotations (many resumes hide the GitHub URL behind the word "GitHub").
pypdf is the fallback. DOCX and TXT are supported as a bonus.

A file that cannot be read raises `UnreadableResume`; the pipeline records
it as failed and carries on with the rest of the batch.
"""

from __future__ import annotations

from pathlib import Path

from . import config
from .models import ParsedResume
from .sections import detect_heading, normalise_text


class UnreadableResume(Exception):
    """The file could not be turned into usable text."""


def _usable(text: str) -> bool:
    compact = "".join(text.split())
    if len(compact) < config.MIN_TEXT_CHARS:
        return False
    letters = sum(ch.isalpha() for ch in compact)
    return letters / len(compact) >= 0.4


def _name_hint_from_spans(page_dict: dict) -> str | None:
    """The largest text on page one is almost always the candidate's name."""
    spans: list[tuple[float, str]] = []
    for block in page_dict.get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = span.get("text", "").strip()
                if text:
                    spans.append((float(span.get("size", 0)), text))
    if not spans:
        return None
    largest = max(size for size, _ in spans)
    words = [text for size, text in spans if size >= largest - 0.5]
    hint = " ".join(" ".join(words).split())
    return hint if 2 <= len(hint) <= 60 else None


def _heading_stats(text: str) -> tuple[int, int]:
    """(headings found, headings with no content before the next heading)."""
    lines = [line for line in normalise_text(text).split("\n") if line.strip()]
    positions = [i for i, line in enumerate(lines) if detect_heading(line)]
    orphans = 0
    for k, position in enumerate(positions):
        following = positions[k + 1] if k + 1 < len(positions) else len(lines)
        if following - position - 1 < 1:
            orphans += 1
    return len(positions), orphans


def _choose_reading_order(content_order: str, position_order: str) -> str:
    """Pick the text order that keeps headings next to their content.

    Content order is right for most files and essential for multi-column
    layouts. Some design tools (Canva, for example) draw all headings
    separately from the body; those files only read correctly when sorted by
    position on the page. Switch only when that clearly fixes stranded
    headings without losing any.
    """
    content_heads, content_orphans = _heading_stats(content_order)
    sorted_heads, sorted_orphans = _heading_stats(position_order)
    if content_orphans >= 1 and sorted_orphans < content_orphans and sorted_heads >= content_heads:
        return position_order
    return content_order


def _parse_pdf_pymupdf(path: Path) -> tuple[str, list[str], int, str | None]:
    import pymupdf

    pymupdf.TOOLS.mupdf_display_errors(False)   # keep font warnings off stderr
    with pymupdf.open(path) as doc:
        if doc.needs_pass:
            raise UnreadableResume("PDF is password protected")
        content_pages: list[str] = []
        sorted_pages: list[str] = []
        links: list[str] = []
        name_hint = None
        for index, page in enumerate(doc):
            content_pages.append(page.get_text("text"))
            sorted_pages.append(page.get_text("text", sort=True))
            for link in page.get_links():
                uri = link.get("uri")
                if uri:
                    links.append(uri.strip())
            if index == 0:
                try:
                    name_hint = _name_hint_from_spans(page.get_text("dict"))
                except Exception:   # a hint is optional
                    name_hint = None
        text = _choose_reading_order("\n".join(content_pages), "\n".join(sorted_pages))
        return text, links, len(content_pages), name_hint


def _parse_pdf_pypdf(path: Path) -> tuple[str, list[str], int, str | None]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise UnreadableResume("PDF is password protected")
    pages: list[str] = []
    links: list[str] = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
        for annotation in page.get("/Annots") or []:
            try:
                action = annotation.get_object().get("/A")
                uri = action.get_object().get("/URI") if action is not None else None
            except Exception:
                uri = None
            if uri:
                links.append(str(uri).strip())
    return "\n".join(pages), links, len(pages), None


def _parse_pdf(path: Path) -> ParsedResume:
    errors: list[str] = []
    for parser_name, parser in (("pymupdf", _parse_pdf_pymupdf), ("pypdf", _parse_pdf_pypdf)):
        try:
            text, links, page_count, name_hint = parser(path)
        except UnreadableResume:
            raise
        except Exception as exc:   # corrupt file, missing library, ...
            errors.append(f"{parser_name}: {type(exc).__name__}: {exc}")
            continue
        text = normalise_text(text)
        if _usable(text):
            warnings = [f"primary parser failed ({errors[0]})"] if errors else []
            return ParsedResume(
                source_file=path.name, path=str(path), file_hash="", text=text,
                links=links, page_count=page_count, parser=parser_name,
                name_hint=name_hint, warnings=warnings,
            )
        errors.append(f"{parser_name}: no usable text (scanned image or empty file?)")
    raise UnreadableResume("; ".join(errors) or "no usable text")


def _parse_docx(path: Path) -> ParsedResume:
    try:
        import docx
    except ImportError as exc:
        raise UnreadableResume("python-docx is not installed; cannot read DOCX") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise UnreadableResume(f"docx: {type(exc).__name__}: {exc}") from exc
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    links = [
        rel.target_ref
        for rel in document.part.rels.values()
        if "hyperlink" in rel.reltype and isinstance(rel.target_ref, str)
    ]
    text = normalise_text("\n".join(parts))
    if not _usable(text):
        raise UnreadableResume("docx: no usable text")
    return ParsedResume(
        source_file=path.name, path=str(path), file_hash="", text=text,
        links=links, page_count=1, parser="python-docx",
    )


def _parse_txt(path: Path) -> ParsedResume:
    text = normalise_text(path.read_text(encoding="utf-8", errors="replace"))
    if not _usable(text):
        raise UnreadableResume("txt: no usable text")
    return ParsedResume(
        source_file=path.name, path=str(path), file_hash="", text=text,
        links=[], page_count=1, parser="text",
    )


def parse_resume(path: Path) -> ParsedResume:
    """Parse one resume. Raises UnreadableResume if nothing usable comes out."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".txt":
        return _parse_txt(path)
    raise UnreadableResume(f"unsupported file type: {suffix}")
